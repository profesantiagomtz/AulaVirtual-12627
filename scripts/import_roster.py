"""Extrae la lista de alumnos del libro de estatus a un JSON privado.

Uso: python scripts/import_roster.py "ruta/al/archivo.xlsx"
El resultado queda en private-data/roster.json, excluido de Git.
"""
from pathlib import Path
import json
import sys
import openpyxl

source = Path(sys.argv[1])
target = Path(__file__).resolve().parents[1] / "private-data" / "roster.json"
workbook = openpyxl.load_workbook(source, data_only=False, read_only=True)
classes = []
for sheet in workbook.worksheets:
    if sheet.title.upper().startswith("TUTORIA"):
        continue
    parts = [part.strip() for part in sheet.title.split("-")]
    subject = parts[0].strip()
    group = parts[-1].strip()
    students = []
    for row in range(4, sheet.max_row + 1):
        number = sheet.cell(row, 1).value
        name = sheet.cell(row, 2).value
        if isinstance(number, (int, float)) and isinstance(name, str) and name.strip():
            students.append({"number": int(number), "name": " ".join(name.split())})
    classes.append({"subject": subject, "group": group, "students": students})
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps({"source": source.name, "classes": classes}, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"output": str(target), "classes": [{"subject": c["subject"], "group": c["group"], "students": len(c["students"])} for c in classes]}, ensure_ascii=False))
