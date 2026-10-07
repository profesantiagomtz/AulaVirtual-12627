from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "public" / "materials" / "edoa"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "12348F"
BLUE = "E4ECFB"
GREEN = "2DA44E"
GRAY = "5B667A"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "D9D9D9")
        element.append(border)
    tbl_pr.append(element)


def base(title, subtitle, number):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph()
    r = p.add_run(f"EDOA  R.A. 1.1  Práctica {number}")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(f"\n{subtitle}")
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    labels = (("Alumno", "Escribe tu nombre completo"), ("Grupo", "311"), ("Fecha", "Escribe la fecha"))
    for i, (label, value) in enumerate(labels):
        shade(table.cell(0, i), NAVY)
        run = table.cell(0, i).paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        table.cell(1, i).text = value
        table.cell(1, i).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return doc


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text)
    p.style.font.color.rgb = RGBColor(0, 0, 0)
    return p


def instruction(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def save(doc, filename):
    doc.core_properties.author = "Aula Virtual"
    doc.core_properties.subject = "Práctica EDOA R.A. 1.1"
    doc.save(OUT / filename)


doc = base("Reconocimiento de Word 2019", "Identifica la función de cada elemento de la interfaz.", 1)
heading(doc, "Instrucciones")
instruction(doc, "Abre Word 2019 y localiza cada elemento de la tabla.")
instruction(doc, "Reemplaza cada texto Escribe aquí con una explicación breve de su función.")
instruction(doc, "No borres los nombres de los elementos ni agregues capturas de otra computadora.")
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(table)
for i, text in enumerate(("Elemento", "Dónde se encuentra", "Para qué sirve")):
    shade(table.cell(0, i), NAVY); table.cell(0, i).text = text
    for run in table.cell(0, i).paragraphs[0].runs: run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
for item in ("Barra de título", "Acceso rápido", "Pestañas", "Cinta de opciones", "Regla", "Área de trabajo", "Barra de estado", "Barra de desplazamiento", "Ayuda"):
    cells = table.add_row().cells
    cells[0].text = item; cells[1].text = "Escribe aquí"; cells[2].text = "Escribe aquí"
heading(doc, "Comprobación")
doc.add_paragraph("Antes de entregar verifica que reemplazaste los 18 espacios Escribe aquí.")
save(doc, "EDOA_RA11_Practica_1_Interfaz.docx")

doc = base("Configuración y diseño de página", "Configura correctamente la hoja antes de dar formato al contenido.", 2)
heading(doc, "Resultado solicitado")
for text in ("Tamaño Carta y orientación vertical.", "Márgenes de 2.54 cm en los cuatro lados.", "Encabezado con tu nombre y grupo.", "Pie de página con el texto Práctica de diseño y número de página.", "Color de página blanco."):
    instruction(doc, text)
heading(doc, "Texto de trabajo")
doc.add_paragraph("La presentación de un documento facilita su lectura y comunica orden. Configurar la página antes de escribir evita cambios inesperados al final del trabajo. Los márgenes, la orientación, el tamaño, el encabezado y el pie de página deben responder al uso que tendrá el documento.")
heading(doc, "Evidencia escrita")
doc.add_paragraph("Escribe aquí una explicación de tres renglones sobre la diferencia entre encabezado y pie de página.")
save(doc, "EDOA_RA11_Practica_2_Diseno_pagina.docx")

doc = base("Formato de texto y párrafo", "Aplica formato sin cambiar las palabras del contenido.", 3)
heading(doc, "Indicaciones")
for text in ("Título: Aptos Display 20, negrita, color azul oscuro y centrado.", "Subtítulo: Aptos 13, cursiva y alineación derecha.", "Primer párrafo: Aptos 11, justificado, interlineado 1.5 y sangría de primera línea de 1.27 cm.", "Segundo párrafo: Calibri 11, alineación izquierda y espacio posterior de 8 puntos.", "Resalta en negrita las palabras claridad, orden y consistencia."):
    instruction(doc, text)
doc.add_paragraph("COMUNICACIÓN DIGITAL EFECTIVA")
doc.add_paragraph("Formato y organización de la información")
doc.add_paragraph("Un documento profesional necesita claridad, orden y consistencia. El formato debe ayudar al lector a reconocer títulos, ideas principales y explicaciones sin distraerlo del contenido.")
doc.add_paragraph("La alineación, la sangría y el interlineado organizan visualmente los párrafos. Antes de entregar, conviene revisar que todos los elementos mantengan un estilo coherente.")
save(doc, "EDOA_RA11_Practica_3_Texto_parrafo.docx")

doc = base("Edición y listas", "Organiza información utilizando edición, viñetas y numeración.", 4)
heading(doc, "Parte A  Ordena los pasos")
doc.add_paragraph("Convierte los siguientes renglones en una lista numerada y colócalos en el orden correcto:")
for text in ("Revisar el documento", "Abrir Word", "Guardar el archivo", "Escribir el contenido", "Configurar la página"):
    doc.add_paragraph(text)
heading(doc, "Parte B  Clasifica los elementos")
doc.add_paragraph("Copia los elementos y crea dos listas con viñetas: Hardware y Software.")
doc.add_paragraph("Teclado  Word  Monitor  Navegador  Impresora  Antivirus  Mouse  Presentaciones")
heading(doc, "Parte C  Buscar y reemplazar")
doc.add_paragraph("Reemplaza todas las apariciones de la palabra archivo por documento en el siguiente texto:")
doc.add_paragraph("El archivo debe guardarse con frecuencia. Antes de compartir el archivo, revisa su nombre. Conserva una copia del archivo en otra ubicación.")
save(doc, "EDOA_RA11_Practica_4_Edicion_listas.docx")

doc = base("Práctica integradora del R.A. 1.1", "Da formato completo al documento y conserva una copia de respaldo.", 5)
heading(doc, "Configuración obligatoria")
for text in ("Tamaño Carta, orientación vertical y márgenes de 2.54 cm.", "Encabezado con nombre, grupo y título del trabajo.", "Número de página arriba a la derecha.", "Fuente Times New Roman 12, interlineado doble y texto justificado.", "Sangría de primera línea de 1.27 cm.", "Aplica estilos Título 1 y Título 2 a los encabezados indicados.", "Guarda el resultado como DOCX y crea una copia en PDF."):
    instruction(doc, text)
doc.add_paragraph("SEGURIDAD DE LA INFORMACIÓN EN LA VIDA DIARIA")
doc.add_paragraph("Introducción")
doc.add_paragraph("La información personal se utiliza diariamente en dispositivos, aplicaciones y servicios digitales. Protegerla requiere hábitos sencillos y decisiones responsables.")
doc.add_paragraph("Recomendaciones")
doc.add_paragraph("Usar contraseñas diferentes. Activar la verificación en dos pasos. Revisar los permisos de las aplicaciones. Evitar enlaces desconocidos. Mantener copias de respaldo.")
doc.add_paragraph("Conclusión")
doc.add_paragraph("Escribe aquí una conclusión personal de cinco renglones sobre la importancia de proteger la información.")
heading(doc, "Control de entrega")
check = doc.add_table(rows=1, cols=2); borders(check); check.style = "Table Grid"
check.cell(0,0).text = "Archivo principal"; check.cell(0,1).text = "Apellido_Nombre_RA11.docx"
row = check.add_row().cells; row[0].text = "Copia de respaldo"; row[1].text = "Apellido_Nombre_RA11.pdf"
save(doc, "EDOA_RA11_Practica_5_Integradora.docx")

print("\n".join(str(path) for path in sorted(OUT.glob("*.docx"))))
